# -*- coding: utf-8 -*-
"""PDF采购订单提取核心逻辑。

特点：
1. 保留固定字段白名单，适合非专业用户勾选。
2. 每条明细保留原始文本块，支持用户自定义“按关键词查找”。
3. 支持当前样例中的 KONE 英文采购订单、KONE/巨人通力中文采购订单、天吴采购单。
4. 导出 Excel 时允许固定字段 + 用户自定义字段，内部字段不会导出。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ALLOWED_COLUMNS = [
    "来源文件", "解析类型", "页码",
    "订单号", "采购单号", "版本号", "制表日期", "制表人", "采购日期", "来源类型", "日期", "修改日期", "打印日期",
    "付款条件", "付款方式", "交货方式", "送货地址", "交货/项目地址", "发票地址", "目的地国家",
    "供应商", "供应商代码", "供应商电话", "供应商传真", "买方联系人", "卖方联系人", "卖方参考号",
    "Buyer VAT No", "买方VAT号", "币种", "货币", "总金额", "税率", "税种", "设备合同号", "账号",
    "项目", "物料", "交货期", "送货日期", "件号", "料号", "品名规格", "物料名称", "物料规格",
    "单位", "数量", "数量原始值", "未税单价", "未税总价", "印刷",
    "Pos", "Material", "Quantity", "Unit", "Price", "Amount", "价格", "金额",
    "Sales order ref", "Sales order ref num", "Sales order ref item", "Sales order no", "Sales order item", "Project ref", "REV", "Shipping instruction",
    "Number of car door entrances", "KCO_RUSH_BUILDING_STATE", "电梯销售地区", "开门方式", "厅门防火等级",
    "包装箱类型", "轿厢外壳内部高度_AK_mm", "轿厢防护板类型",
    "识别号", "图号", "参数及备注", "附加码", "附加信息", "仓库库位", "采购经办人", "审核", "批准",
    "轿厢净开门宽度_LL_mm", "轿门净高_HH_mm", "门尺寸_LL*HH", "高*宽",
    "DIM_CAR_BOX_INNER_LENGTH_mm", "DIM_CAR_BOX_INNER_WIDTH_mm", "DIM_CAR_BOX_INNER_HEIGHT_mm", "长宽高",
]


ALIASES = {
    "文件": "来源文件",
    "文件名": "来源文件",
    "PDF文件": "来源文件",
    "类型": "解析类型",
    "页": "页码",
    "页数": "页码",
    "长*宽*高": "长宽高",
    "L*W*H": "长宽高",
    "尺寸": "长宽高",
    "规格": "物料规格",
    "物料号": "Material",
    "料号": "件号",
    "订单编号": "订单号",
    "采购单号": "订单号",
    "交货日": "送货日期",
    "交货日期": "送货日期",
    "采购量": "数量",
    "数量单位": "数量原始值",
    "原始数量": "数量原始值",
    "Qty": "Quantity",
    "QTY": "Quantity",
    "项目号": "Project ref",
    "项目参考": "Project ref",
    "Project reference": "Project ref",
    "Project Ref": "Project ref",
    "Sales order ref.": "Sales order ref", "Sales order reference": "Sales order ref",
    "Arr.date": "交货期",
    "Req.Shipping.Date": "交货期",
    "Req.Shipping Date": "交货期",
    "Delivery date": "交货期",
    "Delivery Date": "交货期",
    "物料编码": "物料",
    "物料编号": "物料",
    "项目物料": "物料",
    "行号": "项目",
    "序号": "项目",
}


ALIASES.update({
    "Sales order ref no": "Sales order ref num", "Sales order ref number": "Sales order ref num", "Sales order num": "Sales order ref num",
    "Sales order ref line": "Sales order ref item", "Sales order line": "Sales order ref item",
    "DIM_CAR_BOX_INNER_LENGTH": "DIM_CAR_BOX_INNER_LENGTH_mm",
    "DIM_CAR_BOX_INNER_WIDTH": "DIM_CAR_BOX_INNER_WIDTH_mm",
    "DIM_CAR_BOX_INNER_HEIGHT": "DIM_CAR_BOX_INNER_HEIGHT_mm",
    "轿门净高_HH": "轿门净高_HH_mm",
    "轿门净高 HH": "轿门净高_HH_mm",
    "轿门净高....HH": "轿门净高_HH_mm",
    "轿门净高HH": "轿门净高_HH_mm",
    "轿厢净开门宽度_LL": "轿厢净开门宽度_LL_mm",
    "轿厢净开门宽度 LL": "轿厢净开门宽度_LL_mm",
    "轿厢净开门宽度LL": "轿厢净开门宽度_LL_mm",
    "PO号": "订单号", "PO Number": "订单号", "Purchase order No": "订单号", "订单版本": "版本号", "Ver.No.": "版本号",
    "Date changed": "修改日期", "Date printed": "打印日期", "Terms of payment": "付款方式", "Terms of delivery": "交货方式",
    "Delivery address": "交货/项目地址", "Invoicing address": "发票地址", "Country of destination": "目的地国家",
    "Seller/Vendor": "供应商", "Supplier number": "供应商代码", "Seller's contact person": "卖方联系人", "Seller's reference": "卖方参考号", "Buyer's contact person": "买方联系人",
    "买方（甲方） VAT No": "买方VAT号", "VAT No": "买方VAT号", "CURRENCY": "货币", "TOTAL AMOUNT": "总金额", "Total Amount": "总金额", "Currency": "货币", "帐号": "账号",
    "品名": "品名规格", "名称规格": "品名规格", "物料描述": "物料名称", "Description": "物料名称", "REV#": "REV", "Shipping Instruction": "Shipping instruction",
    "Car door entrances": "Number of car door entrances", "轿厢外壳内部高度": "轿厢外壳内部高度_AK_mm", "AK": "轿厢外壳内部高度_AK_mm", "防护板类型": "轿厢防护板类型",
    "参数备注": "参数及备注", "备注": "参数及备注", "仓库": "仓库库位", "库位": "仓库库位",
})

INTERNAL_KEYS = {"_raw_block", "_raw_text", "_source_file", "_parser", "_page"}


TABLE_COLUMN_ALIASES = {
    "项目": "项目",
    "Pos": "Pos",
    "位置": "项目",
    "物料": "物料",
    "物料号": "物料",
    "Material": "Material",
    "交货期": "交货期",
    "送货日期": "送货日期",
    "交货日": "交货期",
    "数量": "数量",
    "采购量": "数量",
    "数量原始值": "数量原始值",
    "Quantity": "Quantity",
    "单位": "单位",
    "Unit": "Unit",
    "PCS": "单位",
    "数量单位": "单位",
    "价格": "价格",
    "单价": "价格",
    "Price": "Price",
    "金额": "金额",
    "总价": "金额",
    "Amount": "Amount",
    "Arr.date": "交货期",
    "Req.Shipping.Date": "交货期",
    "Req.Shipping Date": "交货期",
    "Delivery date": "交货期",
    "图号": "图号",
    "Project ref": "Project ref",
    "Sales order ref num": "Sales order ref num",
    "Sales order ref item": "Sales order ref item",
    "DIM_CAR_BOX_INNER_LENGTH": "DIM_CAR_BOX_INNER_LENGTH_mm",
    "DIM_CAR_BOX_INNER_WIDTH": "DIM_CAR_BOX_INNER_WIDTH_mm",
    "DIM_CAR_BOX_INNER_HEIGHT": "DIM_CAR_BOX_INNER_HEIGHT_mm",
    "轿门净高_HH": "轿门净高_HH_mm",
    "轿门净高 HH": "轿门净高_HH_mm",
    "轿门净高....HH": "轿门净高_HH_mm",
    "轿门净高HH": "轿门净高_HH_mm",
    "轿厢净开门宽度_LL": "轿厢净开门宽度_LL_mm",
    "轿厢净开门宽度 LL": "轿厢净开门宽度_LL_mm",
    "轿厢净开门宽度LL": "轿厢净开门宽度_LL_mm",
}


TABLE_COLUMN_ALIASES.update({
    "采购单号": "订单号", "版本": "版本号", "版本号": "版本号", "日期": "日期",
    "制表日期": "制表日期", "制表人": "制表人", "采购日期": "采购日期", "付款条件": "付款条件", "付款方式": "付款方式", "交货方式": "交货方式",
    "供应商": "供应商", "供应商代码": "供应商代码", "买方联系人": "买方联系人", "货币": "货币", "总金额": "总金额", "税率": "税率", "税种": "税种",
    "料号": "料号", "品名规格": "品名规格", "识别号": "识别号", "参数及备注": "参数及备注", "附加码": "附加码", "附加信息": "附加信息", "仓库库位": "仓库库位",
    "未税单价": "未税单价", "未税金额": "未税总价", "未税总价": "未税总价",
    "REV": "REV", "REV#": "REV", "Shipping instruction": "Shipping instruction", "Number of car door entrances": "Number of car door entrances", "KCO_RUSH_BUILDING_STATE": "KCO_RUSH_BUILDING_STATE",
    "电梯销售地区": "电梯销售地区", "开门方式": "开门方式", "厅门防火等级": "厅门防火等级", "包装箱类型": "包装箱类型", "轿厢外壳内部高度_AK_mm": "轿厢外壳内部高度_AK_mm", "轿厢防护板类型": "轿厢防护板类型",
})

def resolve_table_column_name(name: str) -> str:
    raw = str(name or "").strip()
    if not raw:
        return ""
    if raw in TABLE_COLUMN_ALIASES:
        return TABLE_COLUMN_ALIASES[raw]
    canonical = canonical_column_name(raw)
    return TABLE_COLUMN_ALIASES.get(canonical, canonical)


def canonical_column_name(name: str) -> str:
    return ALIASES.get(str(name or "").strip(), str(name or "").strip())


def clean_text(value: Any) -> str:
    return str(value or "").replace("\xa0", " ").strip()


def normalize_lines(text: str) -> List[str]:
    lines: List[str] = []
    for raw in (text or "").replace("\xa0", " ").replace("\r", "\n").split("\n"):
        line = re.sub(r"[ \t]+", " ", raw).strip()
        if line:
            lines.append(line)
    return lines


def normalize_text(text: str) -> str:
    return "\n".join(normalize_lines(text))


def flatten_text(text: str) -> str:
    return " ".join(normalize_lines(text))


def compact_text(text: str) -> str:
    return re.sub(r"\s+", "", str(text or ""))


def clean_number(value: Any) -> Any:
    if value is None:
        return ""
    text = str(value).strip().replace(",", "")
    if not text:
        return ""
    try:
        number = float(text)
        if number.is_integer():
            return int(number)
        return number
    except Exception:
        return text


def as_plain_number_text(value: Any) -> str:
    value = clean_number(value)
    if value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def split_quantity_unit(value: Any) -> tuple[Any, str, str]:
    raw = str(value or "").strip().replace(" ", "")
    if not raw:
        return "", "", ""
    match = re.match(r"^([\d,]+(?:\.\d+)?)([A-Za-z]+|[\u4e00-\u9fa5]+)?$", raw)
    if not match:
        return raw, "", raw
    return clean_number(match.group(1)), match.group(2) or "", raw


def extract_first_quantity_unit(value: Any) -> tuple[Any, str, str]:
    """从任意文本中提取第一个“数字+单位”片段。

    覆盖常见格式：
    - 1PC / 5.000PCS
    - 4 PC
    - 采购量:5.000PCS
    - 7B 重型瓦楞纸箱,1750X520X300 / 1 PC  （取 / 后的 1 PC 优先）
    """
    text = str(value or "").replace("\xa0", " ").strip()
    if not text:
        return "", "", ""

    # 物料描述里常见“/ 1 PC”，优先取斜杠后的包装单位，避免误取 7B、1750X520X300。
    slash_match = re.search(r"/\s*([0-9,]+(?:\.\d+)?)\s*([A-Za-z]+|[\u4e00-\u9fa5]+)\b", text, re.I)
    if slash_match:
        raw = f"{slash_match.group(1)}{slash_match.group(2)}"
        return clean_number(slash_match.group(1)), slash_match.group(2), raw

    match = re.search(r"(?<![A-Z0-9])([0-9,]+(?:\.\d+)?)\s*([A-Za-z]+|[\u4e00-\u9fa5]+)\b", text, re.I)
    if match:
        raw = f"{match.group(1)}{match.group(2)}"
        return clean_number(match.group(1)), match.group(2), raw

    # 没有单位时，也允许“只取数字”使用第一个数字。
    num_match = re.search(r"(?<![A-Z0-9])([0-9,]+(?:\.\d+)?)", text, re.I)
    if num_match:
        return clean_number(num_match.group(1)), "", num_match.group(1)

    return "", "", ""


def transform_custom_value(value: Any, value_type: str = "raw") -> Any:
    """兼容旧版自定义规则的结果处理。

    v8 起页面不再暴露“只取数字/只取单位”选项；数量和单位统一在底层自动拆分。
    这里仍保留旧配置兼容能力，导入 v7 配置不会失效。
    """
    value_type = str(value_type or "raw").strip()
    if value in ("", None):
        return ""
    if value_type in {"raw", "text", ""}:
        return value

    number, unit, raw = extract_first_quantity_unit(value)
    if value_type in {"number", "quantity"}:
        return number
    if value_type == "unit":
        return unit
    if value_type in {"number_unit", "quantity_unit"}:
        if number == "" and not unit:
            return sanitize_extracted_value(str(value))
        return f"{as_plain_number_text(number)}{unit}" if unit else as_plain_number_text(number)
    return value


def _looks_like_quantity_name(name: Any) -> bool:
    text = str(name or "").strip().lower()
    return any(token in text for token in ["数量", "采购量", "quantity", "qty"])


def _looks_like_unit_name(name: Any) -> bool:
    text = str(name or "").strip().lower()
    return any(token in text for token in ["单位", "unit", "uom"])


def normalise_quantity_unit_from_value(value: Any) -> tuple[Any, str, str]:
    """自动识别常见数量单位：1PC、5.000PCS、4 PC。"""
    if value in ("", None):
        return "", "", ""
    qty, unit, raw = split_quantity_unit(value)
    if qty != "" and unit:
        return qty, unit, raw
    qty, unit, raw = extract_first_quantity_unit(value)
    if qty != "" and unit:
        return qty, unit, raw
    return "", "", ""


def apply_quantity_unit_to_row(row: Dict[str, Any], value: Any, force: bool = False) -> bool:
    """把任意提取值中的数量和单位写入标准字段。

    force=True 用于“采购量/数量/单位”等语义明确字段；普通字段即使含有“/ 1 PC”也不会误覆盖。
    """
    qty, unit, raw = normalise_quantity_unit_from_value(value)
    if qty == "" or not unit:
        return False
    if force or row.get("数量") in ("", None) or row.get("Quantity") in ("", None):
        row["数量"] = qty
        row["Quantity"] = qty
    if force or row.get("单位") in ("", None) or row.get("Unit") in ("", None):
        row["单位"] = unit
        row["Unit"] = unit
    if raw:
        row["数量原始值"] = raw
    return True


def normalize_row_quantity_unit(row: Dict[str, Any]) -> Dict[str, Any]:
    """解析后统一拆分数量和单位。

    目标：页面不再需要“只取数字/只取单位”快捷按钮。只要底层拿到 1PC、5.000PCS、4 PC，
    Excel 选择“数量”时输出数字，选择“单位”时输出单位，选择“数量原始值”时输出原始片段。
    """
    # 先处理语义最明确的字段。
    preferred_keys = ["数量原始值", "采购量", "Quantity raw", "QuantityRaw", "Quantity", "数量"]
    for key in preferred_keys:
        if key in row and row.get(key) not in ("", None):
            if apply_quantity_unit_to_row(row, row.get(key), force=key in {"数量原始值", "采购量", "Quantity raw", "QuantityRaw"}):
                return row

    # 再扫描用户自定义字段名，只有字段名明显表示数量/单位时才自动拆分，避免误取物料规格里的“/ 1 PC”。
    for key, value in list(row.items()):
        if str(key).startswith("_") or value in ("", None):
            continue
        if _looks_like_quantity_name(key) or _looks_like_unit_name(key):
            if apply_quantity_unit_to_row(row, value, force=True):
                return row
    return row


def normalize_rows_quantity_unit(rows: Sequence[Dict[str, Any]]) -> None:
    for row in rows:
        normalize_row_quantity_unit(row)


def normalize_custom_value_for_excel(name: str, title: str, value: Any, row: Dict[str, Any]) -> Any:
    """自定义字段结果写入 Excel 前的自动数量/单位拆分。

    例如用户自定义列名为“采购量”，关键词取到 5.000PCS：
    - 自定义列本身输出 5
    - 系统字段“数量”写 5，“单位”写 PCS，“数量原始值”写 5.000PCS
    """
    if value in ("", None):
        return value
    semantic = f"{name} {title}"
    if _looks_like_quantity_name(semantic) or _looks_like_unit_name(semantic):
        qty, unit, raw = normalise_quantity_unit_from_value(value)
        if qty != "" and unit:
            row["数量"] = qty
            row["Quantity"] = qty
            row["单位"] = unit
            row["Unit"] = unit
            row["数量原始值"] = raw
            if _looks_like_unit_name(semantic) and not _looks_like_quantity_name(semantic):
                return unit
            return qty
    return value


def normalize_sales_order_ref(value: Any) -> str:
    raw = str(value or "").strip()
    if not raw:
        return ""
    raw = raw.replace("／", "/").replace("\\", "/").replace(",", "")
    raw = re.sub(r"\s*/\s*", "/", raw)
    raw = re.sub(r"\s+", "", raw)
    raw = raw.strip(" .:：;，,")
    return raw


def split_sales_order_ref(value: Any) -> tuple[str, str, str]:
    raw = normalize_sales_order_ref(value)
    if not raw:
        return "", "", ""
    parts = raw.split("/", 1)
    if len(parts) == 2:
        return parts[0].strip(), parts[1].strip(), raw
    return raw, "", raw


def find_one(pattern: str, text: str, default: Any = "", flags: int = re.I) -> Any:
    match = re.search(pattern, text or "", flags)
    if not match:
        return default
    return match.group(1).strip()


def find_money_after_label(label: str, text: str) -> Any:
    pattern = rf"{re.escape(label)}\s*[:：]?\s*([0-9,]+(?:\.\d+)?)"
    return clean_number(find_one(pattern, text, flags=re.I))


def format_lwh(length: Any, width: Any, height: Any) -> str:
    values = [as_plain_number_text(length), as_plain_number_text(width), as_plain_number_text(height)]
    if any(v == "" for v in values):
        return ""
    return "*".join(values)


def format_ll_hh(ll: Any, hh: Any) -> str:
    ll_text = as_plain_number_text(ll)
    hh_text = as_plain_number_text(hh)
    if not ll_text or not hh_text:
        return ""
    return f"{ll_text}*{hh_text}"

def format_hh_ll(hh: Any, ll: Any) -> str:
    """按用户要求输出“高*宽”：HH*LL。"""
    hh_text = as_plain_number_text(hh)
    ll_text = as_plain_number_text(ll)
    if not hh_text or not ll_text:
        return ""
    return f"{hh_text}*{ll_text}"


def extract_text_from_pdf(pdf_path: Path) -> str:
    parts: List[str] = []
    with pdfplumber.open(pdf_path) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            parts.append(f"\n---PAGE {page_no}---\n{text}")
    return "\n".join(parts)


def _page_for_line(lines: Sequence[str], index: int) -> int:
    page_no = 1
    for i in range(0, max(index + 1, 0)):
        match = re.match(r"---PAGE\s+(\d+)---", lines[i], re.I)
        if match:
            page_no = int(match.group(1))
    return page_no


def _attach_meta(row: Dict[str, Any], pdf_path: Path, parser: str, page_no: int, block: str, full_text: str) -> Dict[str, Any]:
    row["来源文件"] = pdf_path.name
    row["解析类型"] = parser
    row["页码"] = page_no
    row["_source_file"] = pdf_path.name
    row["_parser"] = parser
    row["_page"] = page_no
    row["_raw_block"] = block
    row["_raw_text"] = full_text
    return row



ORDER_STOP_LABELS = [
    "采购单号", "来源类型", "采购日期", "付款条件", "送货地址", "税率", "税种", "币种", "设备合同号",
    "供应商", "TEL", "FAX", "制表人", "Date", "Date changed", "Date printed", "Terms of delivery",
    "Terms of payment", "Seller's contact person", "Seller's reference", "Buyer's contact person",
    "Supplier number", "CURRENCY", "TOTAL AMOUNT", "采购单号：", "版本号", "电话", "传真", "货币", "总金额",
]


def _remove_label_prefix(value: str, label: str) -> str:
    value = clean_text(value)
    if not value:
        return ""
    lab = re.escape(str(label or "").strip())
    value = re.sub(rf"^{lab}\s*[\.。:：]?\s*", "", value, flags=re.I).strip()
    return value.strip(" ：:;，,")


def find_regex_value(pattern: str, text: str, group: int = 1, flags: int = re.I | re.S, max_len: int = 160) -> str:
    match = re.search(pattern, text or "", flags)
    if not match:
        return ""
    return sanitize_extracted_value(str(match.group(group)), max_len=max_len)


def next_value_after_label_line(text: str, labels: Sequence[str], skip_contains: Sequence[str] | None = None) -> str:
    lines = normalize_lines(text)
    skip = list(skip_contains or [])
    for i, line in enumerate(lines):
        if not any(label.lower() in line.lower() for label in labels):
            continue
        for label in labels:
            m = re.search(rf"{re.escape(label)}\s*[\.。:：]\s*(.+)$", line, re.I)
            if m:
                value = sanitize_extracted_value(m.group(1), max_len=160)
                if value and not any(x.lower() in value.lower() for x in skip):
                    return value
        for nxt in lines[i + 1:i + 5]:
            if not nxt or any(x.lower() in nxt.lower() for x in skip):
                continue
            if any(lbl.lower() in nxt.lower() for lbl in labels):
                continue
            return sanitize_extracted_value(nxt, max_len=160)
    return ""


def extract_labeled_value(text: str, label: str, stop_labels: Sequence[str] | None = None, max_len: int = 180) -> str:
    label = str(label or "").strip()
    if not label:
        return ""
    labels = list(stop_labels or ORDER_STOP_LABELS)
    lines = normalize_lines(text)
    label_compact = compact_text(label).lower()

    for i, line in enumerate(lines):
        line_norm = line.replace("图 号", "图号")
        line_compact = compact_text(line_norm).lower()
        if label_compact not in line_compact:
            continue
        label_for_match = label.replace("图 号", "图号")
        pattern = rf"{re.escape(label_for_match)}\s*[\.。:：]?\s*(.*)$"
        match = re.search(pattern, line_norm, re.I)
        value = match.group(1).strip() if match else ""
        if not value and i + 1 < len(lines):
            next_line = lines[i + 1].strip()
            if next_line and not any(compact_text(next_line).lower().startswith(compact_text(x).lower()) for x in labels):
                value = next_line
        if value:
            for stop in labels:
                if compact_text(stop).lower() == label_compact:
                    continue
                value = re.split(rf"\s+{re.escape(stop)}\s*[\.。:：]?", value, maxsplit=1, flags=re.I)[0]
            value = sanitize_extracted_value(value, max_len=max_len)
            value = _remove_label_prefix(value, label_for_match)
            if value:
                return value
    return ""


def extract_first_date_after_label(text: str, label: str) -> str:
    value = extract_labeled_value(text, label, max_len=80)
    match = re.search(r"\d{4}[/-]\d{1,2}[/-]\d{1,2}|\d{2}\.\d{2}\.\d{4}", value)
    if match:
        return match.group(0)
    return value


def extract_currency_total(text: str) -> tuple[str, Any]:
    flat = flatten_text(text)
    matches = re.findall(r"\b(RMB|CNY|USD|EUR|HKD)\b\s+([0-9,]+(?:\.\d+)?)", flat, flags=re.I)
    if matches:
        currency, amount = matches[-1]
        return currency.upper(), clean_number(amount)
    return "", ""


def extract_block_between(text: str, start_labels: Sequence[str], end_labels: Sequence[str], max_lines: int = 8) -> str:
    lines = normalize_lines(text)
    for i, line in enumerate(lines):
        if not any(lbl.lower() in line.lower() for lbl in start_labels):
            continue
        collected: List[str] = []
        for nxt in lines[i + 1:i + 1 + max_lines]:
            if any(lbl.lower() in nxt.lower() for lbl in end_labels):
                break
            if re.match(r"^\d+\(\s*\d+\)$", nxt):
                break
            collected.append(nxt)
        value = "；".join(collected).strip("； ")
        if value:
            return value[:500]
    return ""


def extract_common_order_fields(full_text: str, parser_name: str = "") -> Dict[str, Any]:
    raw = normalize_text(full_text)
    flat = flatten_text(full_text)
    currency, total_amount = extract_currency_total(full_text)

    fields: Dict[str, Any] = {
        "制表日期": find_regex_value(r"制表日期\s*[:：]\s*([0-9]{4}[/-][0-9]{1,2}[/-][0-9]{1,2})", raw),
        "制表人": find_regex_value(r"制表人\s*[:：]\s*([^\s]+)", raw, max_len=60),
        "采购日期": find_regex_value(r"采购日期\s*[:：]\s*([0-9]{4}[/-][0-9]{1,2}[/-][0-9]{1,2})", raw),
        "来源类型": find_regex_value(r"来源类型\s*[:：]\s*(.*?)(?:\s+采购日期|\s+付款条件|$)", flat, max_len=80),
        "付款条件": find_regex_value(r"付款条件\s*[:：]\s*(.*?)(?:\s+送货地址|\s+税率|$)", flat, max_len=80),
        "送货地址": find_regex_value(r"送货地址\s*[:：]\s*(.*?)(?:\s+税率|\s+税种|\s+币种|\s+设备合同号|$)", flat, max_len=220),
        "税率": find_regex_value(r"税率\s*[:：]\s*([0-9.]+%?)", flat, max_len=40),
        "税种": find_regex_value(r"税种\s*[:：]\s*(.*?)(?:\s+币种|\s+设备合同号|$)", flat, max_len=80),
        "币种": find_regex_value(r"币种\s*[:：]\s*([A-Z]+)", flat, max_len=40),
        "设备合同号": find_regex_value(r"设备合同号\s*[:：]\s*([^\s]+)", flat, max_len=80),
        "供应商": "安徽云博电梯配件有限公司" if "安徽云博电梯配件有限公司" in flat else "",
        "供应商电话": find_regex_value(r"(?:电话|Tel)[:： ]+([0-9+\- ]{5,})", flat, max_len=60),
        "供应商传真": find_regex_value(r"(?:传真|Fax|FAX)[:： ]+([0-9+\- ]{5,})", flat, max_len=60),
        "Buyer VAT No": extract_labeled_value(raw, "Buyer VAT No", max_len=80),
        "买方VAT号": extract_labeled_value(raw, "买方（甲方） VAT No", max_len=80) or extract_labeled_value(raw, "Buyer VAT No", max_len=80),
        "日期": "",
        "修改日期": "",
        "打印日期": "",
        "交货方式": next_value_after_label_line(raw, ["Terms of delivery", "交货方式"], skip_contains=["Seller", "卖方", "Date"]),
        "付款方式": next_value_after_label_line(raw, ["Terms of payment", "付款方式"], skip_contains=["地址", "Road", "发票", "Invoicing"]),
        "卖方联系人": extract_labeled_value(raw, "Seller's contact person", max_len=100) or extract_labeled_value(raw, "卖方联系人", max_len=100),
        "卖方参考号": extract_labeled_value(raw, "Seller's reference", max_len=100) or extract_labeled_value(raw, "卖方参考号", max_len=100),
        "买方联系人": "",
        "供应商代码": "",
        "货币": currency,
        "总金额": total_amount,
        "账号": extract_labeled_value(raw, "帐号", max_len=220) or extract_labeled_value(raw, "账号", max_len=220),
        "交货/项目地址": extract_block_between(raw, ["Delivery address", "交货/项目地址"], ["Country of destination", "目的地国家", "Date", "日期", "Terms"], max_lines=8),
        "发票地址": extract_block_between(raw, ["Invoicing address", "发票地址"], ["Terms of payment", "付款方式", "Buyer's", "买方", "Seller", "卖方"], max_lines=8),
        "目的地国家": extract_labeled_value(raw, "Country of destination", max_len=80) or extract_labeled_value(raw, "目的地国家", max_len=80),
    }

    # 订单号与版本号
    po_match = re.search(r"采购单号\s*[：:]\s*([0-9A-Z_\-]+)", flat, re.I)
    if po_match:
        fields["采购单号"] = po_match.group(1)
        fields["订单号"] = po_match.group(1)
    ver_match = re.search(r"版本号\s*[：:]\s*([0-9A-Z_\-]+)", flat, re.I)
    if ver_match:
        fields["版本号"] = ver_match.group(1)
    po_match = re.search(r"Purchase\s+order\s+No\.\s*([0-9]{6,})", flat, re.I)
    if po_match and not fields.get("订单号"):
        fields["订单号"] = po_match.group(1)
        fields["采购单号"] = po_match.group(1)

    # 日期字段。英文 KONE Date changed / Date printed 常共用同一日期；巨人通力有三行日期。
    dates = re.findall(r"\b\d{2}\.\d{2}\.\d{4}\b", raw)
    if dates:
        fields["日期"] = dates[0]
        if "修改日期" in raw and len(dates) >= 2:
            fields["修改日期"] = dates[1]
        elif "Date changed Date printed" in raw:
            fields["修改日期"] = dates[0]
        if "打印日期" in raw and len(dates) >= 3:
            fields["打印日期"] = dates[2]
        elif "Date changed Date printed" in raw:
            fields["打印日期"] = dates[0]

    contact_code = re.search(r"(?:Buyer.?s contact person\s+Supplier number|买方联系人\s+供应商代码)\s+([A-Za-z,]+)\s+([0-9]{5,})", flat, re.I)
    if contact_code:
        fields["买方联系人"] = contact_code.group(1)
        fields["供应商代码"] = contact_code.group(2)

    pay_match = re.search(r"(?:Terms of payment|付款方式).*?([0-9]{1,3}\s*天[^A-Za-z买方传真Quotation]*)", flat, re.I)
    if pay_match:
        fields["付款方式"] = sanitize_extracted_value(pay_match.group(1), max_len=80)
    elif fields.get("付款方式"):
        pay_text = str(fields["付款方式"])
        pay_match = re.search(r"([0-9]{1,3}\s*天.*)", pay_text)
        fields["付款方式"] = sanitize_extracted_value(pay_match.group(1), max_len=80) if pay_match else ""
    if str(fields.get("版本号", "")).lower() in {"seller", "shipping", "invoicing"}:
        fields["版本号"] = ""

    return {k: v for k, v in fields.items() if v not in ("", None)}


def merge_missing_fields(row: Dict[str, Any], fields: Dict[str, Any]) -> None:
    for key, value in fields.items():
        if value in ("", None):
            continue
        if row.get(key) in ("", None):
            row[key] = value


def extract_ak_height(block: str) -> Any:
    value = find_one(r"轿厢外壳内部高度\s*[,，]?\s*AK\s*\[?mm\]?\s*[:：]?\s*([0-9,]+(?:\.\d+)?)\s*mm", block, flags=re.I)
    if value:
        return clean_number(value)
    compact = compact_text(block)
    value = find_one(r"轿厢外壳内部高度AK(?:MM)?([0-9,]+(?:\.\d+)?)MM?", compact, flags=re.I)
    return clean_number(value) if value else ""

def parse_pdf(pdf_path: Path) -> List[Dict[str, Any]]:
    full_text = extract_text_from_pdf(pdf_path)
    rows: List[Dict[str, Any]] = []

    # 顺序很重要：天吴采购单与 KONE/GiantKONE 差异较大，先解析天吴；KONE 和巨人通力共用表格解析。
    rows.extend(parse_tianwu_pdf(pdf_path, full_text))
    if not rows:
        rows.extend(parse_kone_like_pdf(pdf_path, full_text))
    else:
        # 部分 PDF 理论上只会命中一种格式。保留这里是为了避免误判导致漏提。
        kone_rows = parse_kone_like_pdf(pdf_path, full_text)
        if kone_rows and len(kone_rows) > len(rows):
            rows = kone_rows

    normalize_rows_quantity_unit(rows)
    return rows


def parse_tianwu_pdf(pdf_path: Path, full_text: str) -> List[Dict[str, Any]]:
    raw_text = normalize_text(full_text)
    flat = flatten_text(full_text)
    if not any(key in flat for key in ["采购单号", "料号", "参数及备注", "未税单价", "采购量"]):
        return []

    order_no = find_one(r"采购单号\s*[:：]?\s*([A-Z0-9_\-]+)", flat, flags=re.I)
    if not order_no:
        order_no = find_one(r"\b(PN[0-9A-Z_\-]+)\b", flat, flags=re.I)

    order_fields = extract_common_order_fields(full_text, "天吴采购单")
    supplier_print = extract_printing(flat)
    lines = normalize_lines(full_text)
    starts = [i for i, line in enumerate(lines) if re.search(r"料号\s*[:：]\s*([A-Z0-9_\-]+)", line, re.I)]

    rows: List[Dict[str, Any]] = []
    for idx, line_index in enumerate(starts):
        end_index = starts[idx + 1] if idx + 1 < len(starts) else len(lines)
        block_lines = lines[line_index:end_index]
        block_raw = "\n".join(block_lines)
        block_flat = flatten_text(block_raw)
        item_no = find_one(r"料号\s*[:：]\s*([A-Z0-9_\-]+)", block_flat, flags=re.I)

        qty_raw = find_one(r"采购量\s*[:：]\s*([0-9,]+(?:\.\d+)?\s*[A-Za-z\u4e00-\u9fa5]+)", block_flat, flags=re.I)
        qty, unit, _ = split_quantity_unit(qty_raw)

        remark = find_one(
            r"参数及备注\s*[:：]\s*(.*?)(?:附加码|未税单价|含税单价|采购量|交货日|仓库库位|$)",
            block_flat,
            flags=re.I | re.S,
        )

        material_name = extract_tianwu_material_name(block_flat, remark)
        material_spec = extract_lwh_from_remark(remark) or extract_lwh_from_remark(block_flat)
        length, width, height = split_lwh(material_spec)

        delivery_date = find_one(r"交货日\s*[:：]\s*([0-9]{4}[/-][0-9]{1,2}[/-][0-9]{1,2})", block_flat, flags=re.I)
        unit_price = find_money_after_label("未税单价", block_flat)
        total_amount = find_money_after_label("未税金额", block_flat) or find_money_after_label("未税总价", block_flat)

        row = {
            "项目": idx + 1,
            "物料": item_no,
            "交货期": delivery_date,
            "送货日期": delivery_date,
            "订单号": order_no,
            "采购单号": order_no,
            "件号": item_no,
            "料号": item_no,
            "品名规格": extract_after_label(block_raw, "品名规格") or material_name,
            "物料名称": material_name,
            "物料规格": material_spec,
            "单位": unit,
            "数量": qty,
            "数量原始值": qty_raw,
            "未税单价": unit_price,
            "未税总价": total_amount,
            "印刷": extract_printing(remark) or supplier_print,
            "Material": item_no,
            "Quantity": qty,
            "Unit": unit,
            "Price": unit_price,
            "Amount": total_amount,
            "价格": unit_price,
            "金额": total_amount,
            "DIM_CAR_BOX_INNER_LENGTH_mm": length,
            "DIM_CAR_BOX_INNER_WIDTH_mm": width,
            "DIM_CAR_BOX_INNER_HEIGHT_mm": height,
            "长宽高": material_spec,
            "识别号": extract_after_label(block_raw, "识别号"),
            "图号": extract_after_label(block_raw, "图号") or extract_after_label(block_raw, "图 号"),
            "参数及备注": remark,
            "附加码": extract_after_label(block_raw, "附加码"),
            "附加信息": extract_after_label(block_raw, "附加信息"),
            "仓库库位": extract_after_label(block_raw, "仓库库位"),
        }
        merge_missing_fields(row, order_fields)
        _attach_meta(row, pdf_path, "天吴采购单", _page_for_line(lines, line_index), block_raw, raw_text)
        rows.append(row)

    return rows


def parse_kone_like_pdf(pdf_path: Path, full_text: str) -> List[Dict[str, Any]]:
    raw_text = normalize_text(full_text)
    flat = flatten_text(full_text)
    lines = normalize_lines(full_text)

    # 支持英文 KONE：10 KM52059574V003 23.05.2026 1PC 180.00 180.00
    # 支持中文 GiantKONE：10 KM51950119V004 22.05.2026 4 PC 91.45 365.80
    item_pattern = re.compile(
        r"^(\d{1,5})\s+"
        r"([A-Z0-9]{6,}(?:V\d{3,4})?)\s+"
        r"(\d{2}\.\d{2}\.\d{4})\s+"
        r"([0-9,]+(?:\.\d+)?)\s*"
        r"([A-Za-z]+|[\u4e00-\u9fa5]+)?\s+"
        r"([0-9,]+(?:\.\d{2}))\s+"
        r"([0-9,]+(?:\.\d{2}))\s*$",
        re.I,
    )

    starts: List[tuple[int, re.Match[str]]] = []
    for i, line in enumerate(lines):
        match = item_pattern.match(line)
        if match:
            starts.append((i, match))

    if not starts:
        return []

    po_no = find_one(r"Purchase\s+order\s+No\.\s*([0-9]{6,})", flat, flags=re.I)
    if not po_no:
        po_no = find_one(r"采购单号\s*[：:]\s*([0-9]{6,})", flat, flags=re.I)
    if not po_no:
        po_no = find_one(r"\bNo\.\s*([0-9]{6,})\b", flat, flags=re.I)

    is_giant = "巨人通力" in flat or "GiantKONE" in flat
    parser_name = "巨人通力采购订单" if is_giant else "KONE采购订单"
    order_fields = extract_common_order_fields(full_text, parser_name)

    rows: List[Dict[str, Any]] = []
    for idx, (line_index, match) in enumerate(starts):
        end_index = starts[idx + 1][0] if idx + 1 < len(starts) else len(lines)
        block_lines = lines[line_index:end_index]
        # 去掉每页固定页脚/条款对当前明细自定义字段的干扰，但保留足够上下文。
        block_raw = trim_repeated_footer("\n".join(block_lines))
        block_flat = flatten_text(block_raw)

        qty_text = f"{match.group(4)} {match.group(5) or ''}".strip()
        qty, unit, _ = split_quantity_unit(qty_text)
        sales_ref = extract_sales_order_ref(block_raw) or extract_sales_order_ref(block_flat)
        sales_no, sales_item, _ = split_sales_order_ref(sales_ref)

        hh = extract_hh_mm(block_raw) or extract_hh_mm(block_flat)
        ll = extract_ll_mm(block_raw) or extract_ll_mm(block_flat)

        height = extract_dim_value(block_raw, "HEIGHT") or extract_dim_value(block_flat, "HEIGHT")
        length = extract_dim_value(block_raw, "LENGTH") or extract_dim_value(block_flat, "LENGTH")
        width = extract_dim_value(block_raw, "WIDTH") or extract_dim_value(block_flat, "WIDTH")
        lwh_from_text = extract_lwh_from_remark(block_raw) or extract_lwh_from_remark(block_flat)
        if not (length and width and height) and lwh_from_text:
            length, width, height = split_lwh(lwh_from_text)

        material_code = match.group(2)
        pos_value = clean_number(match.group(1))
        delivery_date = match.group(3)
        price_value = clean_number(match.group(6))
        amount_value = clean_number(match.group(7))

        row = {
            "项目": pos_value,
            "物料": material_code,
            "交货期": delivery_date,
            "送货日期": delivery_date,
            "订单号": po_no,
            "采购单号": po_no,
            "件号": material_code,
            "料号": material_code,
            "物料名称": extract_kone_material_name(block_raw, material_code),
            "品名规格": extract_kone_material_name(block_raw, material_code),
            "物料规格": "",
            "单位": unit,
            "数量": qty,
            "数量原始值": qty_text,
            "未税单价": price_value,
            "未税总价": amount_value,
            "印刷": extract_printing(block_raw),
            "Pos": pos_value,
            "Material": material_code,
            "Quantity": qty,
            "Unit": unit,
            "Price": price_value,
            "Amount": amount_value,
            "价格": price_value,
            "金额": amount_value,
            "Sales order ref": sales_ref,
            "Sales order ref num": sales_no,
            "Sales order ref item": sales_item,
            "Sales order no": sales_no,
            "Sales order item": sales_item,
            "Project ref": extract_after_label(block_raw, "Project ref"),
            "REV": extract_after_label(block_raw, "REV#") or extract_after_label(block_raw, "REV"),
            "Shipping instruction": extract_after_label(block_raw, "Shipping instruction"),
            "Number of car door entrances": extract_after_label(block_raw, "Number of car door entrances"),
            "KCO_RUSH_BUILDING_STATE": extract_after_label(block_raw, "KCO_RUSH_BUILDING_STATE"),
            "电梯销售地区": extract_after_label(block_raw, "电梯销售地区"),
            "开门方式": extract_after_label(block_raw, "开门方式"),
            "厅门防火等级": extract_after_label(block_raw, "厅门防火等级"),
            "图号": extract_after_label(block_raw, "图号"),
            "包装箱类型": extract_after_label(block_raw, "包装箱类型"),
            "轿厢外壳内部高度_AK_mm": extract_ak_height(block_raw),
            "轿厢防护板类型": extract_after_label(block_raw, "轿厢防护板类型"),
            "轿厢净开门宽度_LL_mm": ll,
            "轿门净高_HH_mm": hh,
            "门尺寸_LL*HH": format_ll_hh(ll, hh),
            "高*宽": format_hh_ll(hh, ll),
            "DIM_CAR_BOX_INNER_LENGTH_mm": length,
            "DIM_CAR_BOX_INNER_WIDTH_mm": width,
            "DIM_CAR_BOX_INNER_HEIGHT_mm": height,
            "长宽高": format_lwh(length, width, height),
        }
        row["物料规格"] = row["门尺寸_LL*HH"] or row["长宽高"]
        merge_missing_fields(row, order_fields)
        _attach_meta(row, pdf_path, parser_name, _page_for_line(lines, line_index), block_raw, raw_text)
        rows.append(row)

    return rows


def trim_repeated_footer(block: str) -> str:
    """清理 KONE/GiantKONE 每页重复页脚，但保留跨页续行。

    旧版做法是一遇到页脚标记就把后面的内容全部截断。
    实际采购订单里经常出现“某一条明细从上一页开始，尺寸字段续到下一页页首”的情况；
    直接截断会导致 DIM_CAR_BOX_INNER_WIDTH 等字段丢失。

    新版改为逐行过滤：遇到页脚后仅跳过当前页的固定页脚；遇到下一页页码标记后恢复采集，
    因此可以保留下一页页首属于同一条明细的续行数据。
    """
    lines = normalize_lines(block)
    kept: List[str] = []
    skipping_footer = False

    footer_start_markers = [
        "1. Please acknowledge receipt",
        "For and on behalf of",
        "A KONE Elevators company",
        "帐号:",
        "CURRENCY TOTAL AMOUNT",
        "付款方式",
        "Terms of payment",
        "Seller/Vendor",
        "卖方（乙方）",
        "Invoicing address",
        "发票地址",
        "Delivery address",
        "交货/项目地址",
        "Buyer VAT No",
        "买方（甲方） VAT No",
    ]

    noise_line_patterns = [
        r"^Pos\.\s+Material\s+Arr\.date",
        r"^项目\s+物料\s+交货期",
        r"^KONE Elevators Co\. Ltd\. Purchase order",
        r"^巨人通力电梯有限公司\s+采购订单",
        r"^Purchase order$",
        r"^No\.\s*\d{6,}",
        r"^采购单号[:：]\s*\d{6,}",
        r"^\d+\(\s*\d+\)$",
    ]

    for line in lines:
        # 页码标记来自 extract_text_from_pdf。跨页时，一到新页面就恢复采集，
        # 这样可保留下一页页首的尺寸/项目号等续行。
        if re.match(r"---PAGE\s+\d+---", line, re.I):
            skipping_footer = False
            continue

        if any(marker in line for marker in footer_start_markers):
            skipping_footer = True
            continue

        if skipping_footer:
            continue

        if any(re.search(pattern, line, re.I) for pattern in noise_line_patterns):
            continue

        kept.append(line)

    return "\n".join(kept).strip()


def extract_sales_order_ref(block: str) -> str:
    text = str(block or "")
    if not text:
        return ""

    token = r"[A-Z0-9][A-Z0-9,]{2,30}"
    item_token = r"[A-Z0-9][A-Z0-9,]{0,12}"
    label = r"Sales\s+order\s+ref(?:erence)?\s*\.?\s*[:：]?\s*"

    def is_valid_sales_order_candidate(candidate: str) -> bool:
        candidate = normalize_sales_order_ref(candidate)
        if not re.search(r"\d", candidate):
            return False
        bad_prefixes = ("PROJECTREF", "PURCHASEORDER", "MATERIAL", "DESCRIPTION", "DELIVERY")
        if candidate.upper().startswith(bad_prefixes):
            return False
        return bool(
            re.fullmatch(r"[0-9A-Za-z]+/[0-9A-Za-z]+", candidate)
            or re.fullmatch(r"[0-9A-Za-z]{6,}", candidate)
        )

    patterns = [
        rf"{label}({token}\s*[/／]\s*{item_token})(?=[^A-Z0-9]|$)",
        rf"{label}({token})\s+Sales\s+order\s+item\.?\s*[:：]?\s*({item_token})(?=[^A-Z0-9]|$)",
        rf"{label}({token})\s+(?:item|position|pos\.?|line)\s*[:：]?\s*({item_token})(?=[^A-Z0-9]|$)",
        rf"{label}({token})(?=[^A-Z0-9]|$)",
    ]

    for pattern in patterns:
        match = re.search(pattern, text, re.I)
        if not match:
            continue
        if len(match.groups()) >= 2 and match.group(2):
            candidate = f"{match.group(1)}/{match.group(2)}"
        else:
            candidate = match.group(1)
        candidate = normalize_sales_order_ref(candidate)
        if is_valid_sales_order_candidate(candidate):
            return candidate

    compact = compact_text(text).replace("／", "/")
    match = re.search(r"Salesorderref(?:erence)?\.?[:：]?([A-Z0-9][A-Z0-9,]{2,30}/[0-9][0-9,]{0,12})(?=[^0-9]|$)", compact, re.I)
    if match:
        candidate = normalize_sales_order_ref(match.group(1))
        if is_valid_sales_order_candidate(candidate):
            return candidate

    return ""


def extract_tianwu_material_name(block: str, remark: str) -> str:
    for label in ["品名规格", "物料名称", "品名", "名称"]:
        value = find_one(
            rf"{label}\s*[:：]\s*(.*?)(?:识别号|图号|图 号|参数及备注|未税单价|含税单价|采购量|交货日|$)",
            block,
            flags=re.I | re.S,
        )
        value = clean_material_name(value)
        if value:
            return value

    text = f"{block} {remark}"
    if "板条箱" in text:
        return "板条箱"
    if "木箱" in text and "非加固" in text:
        return "普通木箱"
    if "木箱" in text:
        return "木箱"
    if "纸箱" in text:
        return "纸箱"
    if "托盘" in text:
        return "托盘"
    if "包装箱" in text:
        return "包装箱"
    return ""


def clean_material_name(value: str) -> str:
    value = clean_text(value)
    if not value:
        return ""
    value = re.split(r"(识别号|图号|图 号|参数及备注|未税单价|含税单价|采购量|交货日|仓库库位|REV#)", value)[0]
    value = re.sub(r"[\^;]+", " ", value)
    value = re.sub(r"\s+", " ", value).strip(" ：:;，,")
    return value[:120].strip()


def extract_lwh_from_remark(text: str) -> str:
    if not text:
        return ""

    match = re.search(
        r"\bL\s*=\s*([0-9,]+(?:\.\d+)?)\s*;?\s*W\s*=\s*([0-9,]+(?:\.\d+)?)\s*;?\s*H\s*=\s*([0-9,]+(?:\.\d+)?)",
        text,
        re.I,
    )
    if match:
        return format_lwh(match.group(1), match.group(2), match.group(3))

    # 常见写法：1750X520X300 / 1 PC；2,630*1,170*520。
    match = re.search(r"([0-9,]+(?:\.\d+)?)\s*[*xX×]\s*([0-9,]+(?:\.\d+)?)\s*[*xX×]\s*([0-9,]+(?:\.\d+)?)", text)
    if match:
        return format_lwh(match.group(1), match.group(2), match.group(3))

    return ""


def split_lwh(value: str) -> tuple[Any, Any, Any]:
    if not value:
        return "", "", ""
    parts = re.split(r"\s*[*xX×]\s*", str(value))
    if len(parts) != 3:
        return "", "", ""
    return clean_number(parts[0]), clean_number(parts[1]), clean_number(parts[2])


def extract_printing(text: str) -> str:
    if not text:
        return ""

    match = re.search(r"\^([A-Z]{2,10})\s*;", text)
    if match:
        return match.group(1).strip()

    match = re.search(r"印刷\s*[:：]?\s*([A-Z0-9_\-]+)", text, re.I)
    if match:
        return match.group(1).strip()

    for brand in ["TKE", "KONE", "OTIS", "TK", "蒂升"]:
        if re.search(rf"\b{re.escape(brand)}\b", text, re.I) or brand in text:
            return brand
    return ""


def extract_dim_value(block: str, kind: str) -> Any:
    kind = kind.upper()
    patterns = [
        rf"DIM[_\s]*CAR[_\s]*BOX[_\s]*INNER[_\s]*{kind}(?:[_\s]*mm)?\s*[:：]?\s*([0-9,]+(?:\.\d+)?)\s*(?:mm)?",
        rf"DIM_CAR_BOX_INNER_{kind}\s+([0-9,]+(?:\.\d+)?)\s*mm",
    ]

    for pattern in patterns:
        value = find_one(pattern, block, flags=re.I)
        if value:
            return clean_number(value)

    compact = compact_text(block)
    pattern = rf"DIMCARBOXINNER{kind}(?:MM)?([0-9,]+(?:\.\d+)?)MM?"
    value = find_one(pattern, compact, flags=re.I)
    if value:
        return clean_number(value)

    return ""


def extract_hh_mm(block: str) -> Any:
    patterns = [
        r"轿门净高\.*\s*[,，]?\s*HH\.?\s*\[?mm\]?\s*[:：]?\s*([0-9,]+(?:\.\d+)?)\s*mm",
        r"轿门净高\s*HH\s*\(?mm\)?\s*([0-9,]+(?:\.\d+)?)",
        r"HH\.?\s*\[?mm\]?\s*[:：]?\s*([0-9,]+(?:\.\d+)?)\s*mm",
        r"CAR\s+DOOR\s+HEIGHT.*?([0-9,]+(?:\.\d+)?)\s*mm",
    ]
    for pattern in patterns:
        value = find_one(pattern, block, flags=re.I | re.S)
        if value:
            return clean_number(value)

    compact = compact_text(block)
    value = find_one(r"轿门净高HH(?:MM)?([0-9,]+(?:\.\d+)?)MM?", compact, flags=re.I)
    if value:
        return clean_number(value)
    return ""


def extract_ll_mm(block: str) -> Any:
    patterns = [
        r"轿厢净开门宽度\s*[,，]?\s*LL\.?\s*\(?\[?mm\]?\)?\s*[:：]?\s*([0-9,]+(?:\.\d+)?)\s*mm",
        r"轿厢净开门宽度\s*LL\s*\(?mm\)?\s*([0-9,]+(?:\.\d+)?)",
        r"LL\.?\s*\[?mm\]?\s*[:：]?\s*([0-9,]+(?:\.\d+)?)\s*mm",
        r"CAR\s+DOOR\s+WIDTH.*?([0-9,]+(?:\.\d+)?)\s*mm",
    ]
    for pattern in patterns:
        value = find_one(pattern, block, flags=re.I | re.S)
        if value:
            return clean_number(value)

    compact = compact_text(block)
    value = find_one(r"轿厢净开门宽度LL(?:MM)?([0-9,]+(?:\.\d+)?)MM?", compact, flags=re.I)
    if value:
        return clean_number(value)
    return ""


def extract_kone_material_name(block: str, material_code: str) -> str:
    # 优先保留明细行下一行的真实描述，避免只输出“纸箱/木箱”而丢掉 Q 值、7B、带垫板等信息。
    desc = extract_kone_description_line(block, material_code)
    if desc:
        return desc

    for label in ["Material description", "Item description", "Description"]:
        value = find_one(rf"{label}\s*[:：]?\s*([^\n]+)", block, flags=re.I)
        value = clean_material_name(value)
        if value:
            return value

    if "板条箱" in block:
        return "板条箱"
    if "胶合板箱" in block:
        return "胶合板箱"
    if "木箱" in block and "非加固" in block:
        return "普通木箱"
    if "木箱" in block:
        return "木箱"
    if "纸箱" in block or "瓦楞纸箱" in block:
        return "纸箱"
    if "包装箱" in block:
        return "包装箱"

    idx = block.find(material_code)
    if idx >= 0:
        tail = block[idx + len(material_code): idx + len(material_code) + 220]
        tail = re.sub(r"\d{2}\.\d{2}\.\d{4}.*", "", tail).strip()
        tail = clean_material_name(tail)
        if tail and not re.fullmatch(r"[0-9A-Za-z .,\-/]+", tail):
            return tail

    return ""


def extract_kone_description_line(block: str, material_code: str) -> str:
    lines = normalize_lines(block)
    if not lines:
        return ""

    item_line_index = -1
    for i, line in enumerate(lines):
        if material_code in line and re.search(r"\d{2}\.\d{2}\.\d{4}", line):
            item_line_index = i
            break
    if item_line_index < 0:
        item_line_index = 0

    stop_prefixes = (
        "Sales order ref", "Project ref", "DIM_CAR_BOX", "Number of", "轿门", "轿厢",
        "开门方式", "厅门", "Shipping instruction", "REV#", "图号", "包装箱类型",
    )
    for line in lines[item_line_index + 1:item_line_index + 5]:
        if not line or any(line.startswith(prefix) for prefix in stop_prefixes):
            continue
        if re.match(r"^\d{1,5}\s+[A-Z0-9]{6,}", line, re.I):
            continue
        # 典型描述：胶合板箱 / 1 PC、板条箱,Q=5 / 1 PC、7B 重型瓦楞纸箱,1750X520X300 带垫板 / 1 PC。
        if any(key in line for key in ["箱", "板", "托盘", "crate", "carton", "case"]):
            value = re.sub(r"\s*/\s*\d+(?:\.\d+)?\s*[A-Za-z]+\s*$", "", line, flags=re.I)
            value = re.sub(r"\s+", " ", value).strip(" ：:;，,")
            return value[:120]
    return ""


def extract_after_label(text: str, label: str) -> str:
    label = str(label or "").strip()
    if not label:
        return ""
    # 逐行查找更适合普通用户理解：关键词后面的内容，到本行结束。
    for line in normalize_lines(text):
        # 兼容 “图 号” 与 “图号”。
        line_for_match = line.replace("图 号", "图号")
        label_for_match = label.replace("图 号", "图号")
        if label_for_match.lower() not in line_for_match.lower():
            continue
        pattern = rf"{re.escape(label_for_match)}\s*[\.。:：]?\s*(.*)$"
        match = re.search(pattern, line_for_match, re.I)
        if match:
            value = match.group(1).strip(" ：:;，,")
            return sanitize_extracted_value(value)
    return ""


def sanitize_extracted_value(value: str, max_len: int = 160) -> str:
    value = clean_text(value)
    # 避免把明显的下一字段吞进去。
    value = re.split(
        r"\s+(?:Sales\s+order\s+ref|Project\s+ref|DIM_CAR_BOX|REV#|Shipping\s+instruction|图号|采购量|交货日|未税单价|原材料仓|有图外购件)\b",
        value,
        flags=re.I,
    )[0]
    return value[:max_len].strip(" ：:;，,")


def extract_custom_value(text: str, rule: Dict[str, Any]) -> Any:
    mode = str(rule.get("mode") or "after_label").strip()
    label = str(rule.get("label") or "").strip()
    end_label = str(rule.get("end_label") or "").strip()
    value_type = str(rule.get("value_type") or "raw").strip()
    if not label and mode != "regex":
        return ""

    if mode == "after_label":
        return transform_custom_value(extract_after_label(text, label), value_type)

    if mode == "next_line":
        lines = normalize_lines(text)
        for i, line in enumerate(lines):
            if label.lower() in line.lower() and i + 1 < len(lines):
                return transform_custom_value(sanitize_extracted_value(lines[i + 1]), value_type)
        return ""

    if mode == "between_labels":
        if not end_label:
            return ""
        pattern = rf"{re.escape(label)}\s*[\.。:：]?\s*(.*?)\s*{re.escape(end_label)}"
        match = re.search(pattern, text or "", re.I | re.S)
        if match:
            return transform_custom_value(sanitize_extracted_value(re.sub(r"\s+", " ", match.group(1))), value_type)
        return ""

    # 表格列模式需要读取当前 row 中已解析好的字段，在 apply_custom_fields 中处理。
    if mode == "table_column":
        return ""

    # 高级模式，仅给确实需要的用户使用；前端默认不展示。
    if mode == "regex":
        pattern = str(rule.get("pattern") or "").strip()
        if not pattern:
            return ""
        try:
            match = re.search(pattern, text or "", re.I | re.S)
        except re.error:
            return ""
        if not match:
            return ""
        value = match.group(1) if match.groups() else match.group(0)
        return transform_custom_value(sanitize_extracted_value(re.sub(r"\s+", " ", value)), value_type)

    return ""


def apply_custom_fields(rows: List[Dict[str, Any]], custom_fields: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    for row in rows:
        item_text = str(row.get("_raw_block") or "")
        order_text = str(row.get("_raw_text") or item_text)
        for rule in custom_fields:
            name = str(rule.get("name") or "").strip()
            if not name or name.startswith("_"):
                continue

            mode = str(rule.get("mode") or "after_label").strip()

            # 自定义查找项中的“表格列值”模式：
            # 用户添加“项目、物料、交货期、数量、价格、金额”时，不再从文本关键词后面找，
            # 而是直接读取解析器已经识别好的当前明细列。
            # 这能覆盖 GiantKONE 中文表头：项目 物料 交货期 数量 价格 金额。
            if mode == "table_column":
                source_name = resolve_table_column_name(str(rule.get("label") or name))
                value = get_row_value(row, source_name)
                if value in ("", None) and source_name != name:
                    value = get_row_value(row, name)
                value = transform_custom_value(value, str(rule.get("value_type") or "raw"))
                row[name] = normalize_custom_value_for_excel(name, str(rule.get("title") or name), value, row)
                continue

            scope = str(rule.get("scope") or "item").strip()
            source_text = order_text if scope == "order" else item_text
            value = extract_custom_value(source_text, rule)

            # 表格列（如“项目、物料、交货期、数量、价格、金额”）是系统字段。
            # 如果用户仍然用“关键词后取值”添加同名字段，关键词方式可能找不到值；
            # 这里自动回退到系统列，避免预览里显示 0/N 行找到。
            if value in ("", None):
                fallback_name = resolve_table_column_name(str(rule.get("label") or name))
                fallback_value = get_row_value(row, fallback_name)
                if fallback_value not in ("", None):
                    value = transform_custom_value(fallback_value, str(rule.get("value_type") or "raw"))

            if value not in ("", None) or name not in row:
                row[name] = normalize_custom_value_for_excel(name, str(rule.get("title") or name), value, row)
    normalize_rows_quantity_unit(rows)
    return rows


def get_row_value(row: Dict[str, Any], column: str) -> Any:
    if not column or str(column).startswith("_"):
        return ""
    canonical = canonical_column_name(column)
    if canonical in row:
        return row.get(canonical, "")
    return row.get(str(column).strip(), "")


def _normalise_excel_columns(rows: Sequence[Dict[str, Any]], columns: Iterable[Any]) -> List[Dict[str, str]]:
    """把导出列统一为 {key, title}。

    兼容旧版字符串列表，也支持新版：
    [{"key": "Material", "title": "物料", "order": 2}, ...]
    """
    result: List[Dict[str, str]] = []
    row_keys = set()
    for row in rows:
        row_keys.update(k for k in row.keys() if not str(k).startswith("_"))

    used_pairs: set[tuple[str, str]] = set()
    for item in columns:
        if isinstance(item, dict):
            raw_key = str(item.get("key") or item.get("name") or item.get("title") or "").strip()
            title = str(item.get("title") or raw_key).strip()
        else:
            raw_key = str(item or "").strip()
            title = raw_key

        key = canonical_column_name(raw_key)
        if not key or key.startswith("_") or not title:
            continue

        # 固定白名单字段始终允许；用户自定义字段只要已经存在于行数据中，也允许导出。
        if key not in ALLOWED_COLUMNS and key not in row_keys:
            continue

        pair = (key, title)
        if pair in used_pairs:
            continue
        used_pairs.add(pair)
        result.append({"key": key, "title": title})
    return result


def write_excel(rows: Sequence[Dict[str, Any]], columns: Iterable[Any], output_path: Path, errors: Sequence[Dict[str, Any]] | None = None) -> None:
    errors = list(errors or [])
    rows = list(rows or [])
    column_defs = _normalise_excel_columns(rows, columns)

    wb = Workbook()
    ws = wb.active
    ws.title = "提取结果"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_index, col in enumerate(column_defs, start=1):
        cell = ws.cell(row=1, column=col_index, value=col["title"])
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, row in enumerate(rows, start=2):
        for col_index, col in enumerate(column_defs, start=1):
            value = get_row_value(row, col["key"])
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    if rows and column_defs:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(column_defs))}{len(rows) + 1}"

    for col_index, col in enumerate(column_defs, start=1):
        max_len = len(str(col["title"]))
        for row_index in range(2, min(len(rows) + 2, 300)):
            value = ws.cell(row=row_index, column=col_index).value
            if value is not None:
                max_len = max(max_len, min(len(str(value)), 60))
        ws.column_dimensions[get_column_letter(col_index)].width = max(10, min(max_len + 4, 60))

    money_keys = {"未税单价", "未税总价", "Price", "Amount", "价格", "金额"}
    money_titles = {"单价", "总价", "价格", "金额", "未税单价", "未税金额", "未税总价"}
    for col_index, col in enumerate(column_defs, start=1):
        if col["key"] in money_keys or col["title"] in money_titles:
            for row_index in range(2, len(rows) + 2):
                ws.cell(row=row_index, column=col_index).number_format = "#,##0.00"

    log_ws = wb.create_sheet("错误日志")
    log_headers = ["PDF文件", "状态/错误"]
    for col_index, name in enumerate(log_headers, start=1):
        cell = log_ws.cell(row=1, column=col_index, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border

    if errors:
        for row_index, error in enumerate(errors, start=2):
            log_ws.cell(row=row_index, column=1, value=error.get("file", "")).border = border
            log_ws.cell(row=row_index, column=2, value=error.get("error", "")).border = border
    else:
        log_ws.cell(row=2, column=1, value="无").border = border
        log_ws.cell(row=2, column=2, value="全部PDF处理成功").border = border

    log_ws.column_dimensions["A"].width = 45
    log_ws.column_dimensions["B"].width = 100

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
